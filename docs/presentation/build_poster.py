"""
Build C-Day Computing Showcase poster for GexVisor.
Uses the official KSU c-day-template4.pptx as base, preserving branding.
PRINT-FRIENDLY: white/light background, dark text, minimal ink coverage.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(SCRIPT_DIR, "c-day-template4.pptx")
OUTPUT = os.path.join(SCRIPT_DIR, "GexVisor_CDay_Poster.pptx")

# ── Print-friendly color palette ──────────────────────────────────────
# Template's footer/header are already dark — we keep those as branding.
# Everything else: white bg, dark text, gold accents only where needed.
KSU_GOLD = RGBColor(0xFF, 0xC0, 0x00)
KSU_DARK_GOLD = RGBColor(0xC4, 0x8E, 0x00)  # Readable gold for light bg
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
BODY_TEXT = RGBColor(0x2D, 0x2D, 0x2D)
LIGHT_BG = RGBColor(0xF7, 0xF7, 0xF7)       # Very light gray content bg
PLACEHOLDER_BG = RGBColor(0xEC, 0xEC, 0xEC)  # Placeholder boxes
PLACEHOLDER_BORDER = RGBColor(0xAA, 0xAA, 0xAA)
SECTION_BAR = RGBColor(0x1A, 0x1A, 0x1A)     # Black section header bar
SECTION_TEXT = WHITE                           # White text on black bar
LIGHT_LINE = RGBColor(0xCC, 0xCC, 0xCC)

prs = Presentation(TEMPLATE)
slide = prs.slides[0]

# ── Identify shapes to keep vs remove ──────────────────────────────────
KEEP_NAMES = {'object 28', 'Picture 22', 'Straight Connector 24', 'Picture 13'}
KEEP_TITLE = 'object 2'
KEEP_AUTHOR = 'object 39'
KEEP_NUMBER = 'Text Box 19'

shapes_to_remove = []
for shape in slide.shapes:
    if shape.name not in KEEP_NAMES and shape.name != KEEP_TITLE \
       and shape.name != KEEP_AUTHOR and shape.name != KEEP_NUMBER:
        shapes_to_remove.append(shape)

for shape in shapes_to_remove:
    sp = shape._element
    sp.getparent().remove(sp)


# ── Update title ───────────────────────────────────────────────────────
for shape in slide.shapes:
    if shape.name == KEEP_TITLE:
        tf = shape.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = "GexVisor: A Cross-Platform Framework for Interactive Gamma Exposure Analysis"
        run.font.name = "Arial"
        run.font.size = Pt(110)
        run.font.bold = True
        run.font.color.rgb = WHITE

    elif shape.name == KEEP_NUMBER:
        tf = shape.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = "CS-PhD"
        run.font.size = Pt(120)
        run.font.bold = True
        run.font.color.rgb = WHITE

    elif shape.name == KEEP_AUTHOR:
        tf = shape.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = "Christopher Regan"
        run.font.name = "Arial"
        run.font.size = Pt(60)
        run.font.bold = True
        run.font.color.rgb = KSU_GOLD
        p2 = tf.add_paragraph()
        run2 = p2.add_run()
        run2.text = "Advisor: Dr. Ying Xie"
        run2.font.name = "Arial"
        run2.font.size = Pt(60)
        run2.font.bold = True
        run2.font.color.rgb = KSU_GOLD
        p3 = tf.add_paragraph()
        run3 = p3.add_run()
        run3.text = "College of Computing and Software Engineering"
        run3.font.name = "Arial"
        run3.font.size = Pt(44)
        run3.font.bold = False
        run3.font.color.rgb = WHITE


# ── Helper functions ───────────────────────────────────────────────────
def add_rect(left, top, width, height, fill_color=None, line_color=None, line_width=Pt(0)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.shadow.inherit = False
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = line_width
    else:
        shape.line.fill.background()
    return shape


def add_rounded_rect(left, top, width, height, fill_color=None, line_color=None, line_width=Pt(0)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.shadow.inherit = False
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = line_width
    else:
        shape.line.fill.background()
    return shape


def add_text_box(left, top, width, height, text, font_size=Pt(40),
                 font_color=BODY_TEXT, bold=False, alignment=PP_ALIGN.LEFT,
                 font_name="Arial"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    txBox.text_frame.word_wrap = True
    txBox.text_frame.auto_size = None
    p = txBox.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = font_size
    run.font.color.rgb = font_color
    run.font.bold = bold
    run.font.name = font_name
    p.alignment = alignment
    return txBox


def add_image_placeholder(left, top, width, height, label, sublabel=""):
    """Light gray box with border — drop screenshot here later."""
    shape = add_rounded_rect(left, top, width, height,
                             fill_color=PLACEHOLDER_BG,
                             line_color=PLACEHOLDER_BORDER, line_width=Pt(1.5))
    shape.line.dash_style = 4  # Dash

    add_text_box(left, top + height // 2 - Inches(0.45), width, Inches(0.6),
                 label, font_size=Pt(36), font_color=PLACEHOLDER_BORDER,
                 bold=True, alignment=PP_ALIGN.CENTER)
    if sublabel:
        add_text_box(left, top + height // 2 + Inches(0.2), width, Inches(0.5),
                     sublabel, font_size=Pt(22), font_color=PLACEHOLDER_BORDER,
                     alignment=PP_ALIGN.CENTER)
    return shape


def add_section_header(left, top, width, text):
    """Black bar with gold accent stripe and white text — compact, print-friendly."""
    # Black bar background
    add_rect(left, top, width, Inches(0.75), fill_color=SECTION_BAR)
    # Gold accent left edge
    add_rect(left, top, Inches(0.2), Inches(0.75), fill_color=KSU_GOLD)
    # White text on the bar
    add_text_box(left + Inches(0.4), top + Inches(0.05), width - Inches(0.5), Inches(0.65),
                 text, font_size=Pt(44), font_color=SECTION_TEXT,
                 bold=True, font_name="Arial")
    return top + Inches(0.9)


def add_bullet_text(left, top, width, items, font_size=Pt(32), font_color=BODY_TEXT):
    """Add a text box with bullet points — dark text on light bg."""
    txBox = slide.shapes.add_textbox(left, top, width, Inches(len(items) * 0.6 + 0.1))
    txBox.text_frame.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = txBox.text_frame.paragraphs[0]
        else:
            p = txBox.text_frame.add_paragraph()
        run = p.add_run()
        run.text = item
        run.font.size = font_size
        run.font.color.rgb = font_color
        run.font.name = "Arial"
        p.space_after = Pt(6)
    return txBox


# ═══════════════════════════════════════════════════════════════════════
# WHITE CONTENT BACKGROUND — print-friendly
# Template keeps its black header bar and gold footer natively.
# We fill the main content area with white.
# ═══════════════════════════════════════════════════════════════════════

# White background for entire content area (below title, above footer)
add_rect(Inches(0), Inches(3.0), Inches(48), Inches(30.0), fill_color=WHITE)

# Light gray right panel to visually distinguish screenshot area
add_rect(Inches(14.75), Inches(3.0), Inches(33.25), Inches(30.0),
         fill_color=LIGHT_BG)

# Thin gold separator line between left and right
add_rect(Inches(14.5), Inches(3.5), Inches(0.08), Inches(29.0), fill_color=KSU_GOLD)

# ── LEFT COLUMN (x=1.5 to 13.5) ──────────────────────────────────────
LEFT_X = Inches(1.5)
LEFT_W = Inches(12.0)
y = Inches(3.5)

# Subtitle / engineering angle
add_text_box(LEFT_X, y, LEFT_W, Inches(2.0),
             "Architecture, WebAssembly, and Test Infrastructure "
             "for a .NET Financial Visualization Platform",
             font_size=Pt(44), font_color=KSU_DARK_GOLD, bold=True,
             alignment=PP_ALIGN.LEFT)
y += Inches(2.5)

# The Problem
y = add_section_header(LEFT_X, y, LEFT_W, "The Problem")
add_bullet_text(LEFT_X + Inches(0.2), y, LEFT_W - Inches(0.4), [
    "\u2022 Existing GEX tools: fragmented JS scripts, no tests, no CI",
    "\u2022 Options visualization needs real-time + historical views",
    "\u2022 Research workflows disconnected from analysis tooling",
    "\u2022 No cross-platform solution (browser-based needed)",
], font_size=Pt(34))
y += Inches(3.2)

# Architecture
y = add_section_header(LEFT_X, y, LEFT_W, "System Architecture")
add_image_placeholder(LEFT_X, y, LEFT_W, Inches(6.0),
                      "ARCHITECTURE DIAGRAM",
                      "3-tier: UI \u2192 Core \u2192 API")
y += Inches(6.3)

# Tech stack
y = add_section_header(LEFT_X, y, LEFT_W, "Technology Stack")
add_bullet_text(LEFT_X + Inches(0.2), y, LEFT_W - Inches(0.4), [
    "\u2022 .NET 10 / Blazor WebAssembly (runs in browser via WASM)",
    "\u2022 C# throughout \u2014 replaced JS/Python prototype entirely",
    "\u2022 ApexCharts for interactive financial visualizations",
    "\u2022 SQLite caching layer for market data persistence",
    "\u2022 ASP.NET Core API with options chain services",
    "\u2022 GitHub Actions CI/CD + Husky.Net pre-commit hooks",
], font_size=Pt(32))

# ── RIGHT AREA — split into 2 columns of screenshots ─────────────────
R_LEFT = Inches(15.5)
R_WIDTH = Inches(31.0)
COL_GAP = Inches(0.6)
COL_W = (R_WIDTH - COL_GAP) / 2
col_a = R_LEFT
col_b = R_LEFT + COL_W + COL_GAP

# ─── Right Column A ──────────────────────────────────────────────────
y = Inches(3.5)

y = add_section_header(col_a, y, COL_W, "GEX Visualizer")
add_image_placeholder(col_a, y, COL_W, Inches(6.5),
                      "GEX VISUALIZER",
                      "Interactive chart + regime timeline + annotations")
y += Inches(6.9)

y = add_section_header(col_a, y, COL_W, "Research Notebook")
add_image_placeholder(col_a, y, COL_W, Inches(6.0),
                      "RESEARCH NOTEBOOK",
                      "KDD process tracking + complexity radar")
y += Inches(6.4)

y = add_section_header(col_a, y, COL_W, "Paper Trading Journal")
add_image_placeholder(col_a, y, COL_W, Inches(6.0),
                      "PAPER TRADING",
                      "Trade grid + decision timeline + regime context")
y += Inches(6.4)

# Stats caption
add_text_box(col_a, y, COL_W, Inches(0.6),
             "15 pages  \u2022  35+ Blazor components  \u2022  Full keyboard navigation",
             font_size=Pt(28), font_color=KSU_DARK_GOLD,
             bold=True, alignment=PP_ALIGN.CENTER)

# ─── Right Column B ──────────────────────────────────────────────────
y = Inches(3.5)

y = add_section_header(col_b, y, COL_W, "Research Arcade")
add_image_placeholder(col_b, y, COL_W, Inches(5.5),
                      "RESEARCH ARCADE",
                      "Pattern discovery + data pipeline state")
y += Inches(5.9)

y = add_section_header(col_b, y, COL_W, "Task Board & CI/CD")
half_w = (COL_W - Inches(0.4)) / 2
add_image_placeholder(col_b, y, half_w, Inches(5.5),
                      "TASK BOARD", "GitHub Kanban")
add_image_placeholder(col_b + half_w + Inches(0.4), y, half_w, Inches(5.5),
                      "CI/CD PIPELINE", "GitHub Actions")
y += Inches(5.9)

y = add_section_header(col_b, y, COL_W, "Test Infrastructure")
add_image_placeholder(col_b, y, half_w, Inches(5.0),
                      "TEST RESULTS", "xUnit + Moq + Coverlet")
add_image_placeholder(col_b + half_w + Inches(0.4), y, half_w, Inches(5.0),
                      "BACKTEST TRACKER", "Strategy validation")
y += Inches(5.4)

# Engineering decisions
y = add_section_header(col_b, y, COL_W, "Key Engineering Decisions")
add_bullet_text(col_b + Inches(0.2), y, COL_W - Inches(0.4), [
    "\u2022 JS/Python \u2192 .NET: type safety, single-language stack",
    "\u2022 WebAssembly: full app runs client-side in browser",
    "\u2022 3-tier separation: UI / Core / API testable independently",
    "\u2022 Load testing suite for API performance validation",
    "\u2022 StyleCop + EditorConfig for consistent code quality",
], font_size=Pt(30))

# ── QR Code placeholder (above footer, right side) ────────────────────
add_image_placeholder(Inches(43.0), Inches(27.5), Inches(3.5), Inches(3.5),
                      "QR CODE", "Link to repo/demo")

# ── Save ───────────────────────────────────────────────────────────────
prs.save(OUTPUT)
print(f"Poster saved to: {OUTPUT}")
